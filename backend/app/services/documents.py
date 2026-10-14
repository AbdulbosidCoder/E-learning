"""Turn an uploaded file into something we can send to Claude."""
import io
import os
from dataclasses import dataclass

MAX_PDF_BYTES = 30 * 1024 * 1024


class UnsupportedDocument(ValueError):
    pass


@dataclass
class SourceDocument:
    filename: str
    text: str | None = None  # for DOCX / TXT / MD
    pdf: bytes | None = None  # PDFs are sent to Claude as-is, so scanned pages and tables survive


def _docx_text(raw: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(raw))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


def load(filename: str, raw: bytes) -> SourceDocument:
    ext = os.path.splitext(filename)[1].lower()
    if not raw:
        raise UnsupportedDocument("File is empty")
    if ext == ".pdf":
        if len(raw) > MAX_PDF_BYTES:
            raise UnsupportedDocument("PDF is larger than 30 MB")
        return SourceDocument(filename=filename, pdf=raw)
    if ext == ".docx":
        text = _docx_text(raw)
    elif ext in (".txt", ".md"):
        text = raw.decode("utf-8", errors="replace")
    else:
        raise UnsupportedDocument("Supported formats: PDF, DOCX, TXT, MD")
    if not text.strip():
        raise UnsupportedDocument("No text found in the document")
    return SourceDocument(filename=filename, text=text)
